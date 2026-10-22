import io
import openpyxl
from django.test import TestCase, Client
from django.core.files.uploadedfile import SimpleUploadedFile
from organizations.models import Organization, SuppressionEntry
from core.models import User
from contacts.models import Contact, Tag

class ContactImportTestCase(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(
            name="Org Test Import",
            trade_name="Test Import",
            email="import@test.com",
            max_contacts=100
        )
        self.admin_user = User.objects.create_user(
            username="admin_import",
            email="admin@test.com",
            password="pass1234_password",
            role="org_admin",
            organization=self.org
        )
        self.editor_user = User.objects.create_user(
            username="editor_import",
            email="editor@test.com",
            password="pass1234_password",
            role="campaign_editor",
            organization=self.org
        )
        self.client = Client()

    def test_import_excel_with_only_email_and_nombre(self):
        """Prueba importación de archivo Excel que solo tiene columnas 'email' y 'nombre' (caso usuario)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hoja1"
        ws.append(["email", "nombre"])
        ws.append(["sandra.salazar.agronoma@gmail.com", "Sandra"])
        ws.append(["danirivera2017@udec.cl", "Daniel"])
        ws.append(["noemicatrilaf3@gmail.com", "Noemí"])
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        excel_file = SimpleUploadedFile(
            "Contactos_Prueba.xlsx",
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
        self.client.force_login(self.admin_user)
        response = self.client.post('/contacts/import/', {'file': excel_file, 'tag': 'Clientes 2026'}, follow=True)
        self.assertEqual(response.status_code, 200)
        
        # Verificar contactos creados
        self.assertEqual(Contact.objects.filter(organization=self.org).count(), 3)
        sandra = Contact.objects.get(organization=self.org, email="sandra.salazar.agronoma@gmail.com")
        self.assertEqual(sandra.first_name, "Sandra")
        self.assertEqual(sandra.last_name, "")
        self.assertEqual(sandra.company, "")
        self.assertEqual(sandra.phone, "")
        self.assertTrue(sandra.tags.filter(name="Clientes 2026").exists())

    def test_import_excel_with_duplicate_emails(self):
        """Verifica que filas con emails duplicados no provoquen errores y actualicen la información."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["email", "nombre"])
        ws.append(["dendroenergiapromotor@gmail.com", "Rodrigo"])
        ws.append(["dendroenergiapromotor@gmail.com", "Juan"])
        ws.append(["dendroenergiapromotor@gmail.com", "Luis"])
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        excel_file = SimpleUploadedFile("duplicados.xlsx", buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.client.force_login(self.admin_user)
        response = self.client.post('/contacts/import/', {'file': excel_file}, follow=True)
        self.assertEqual(response.status_code, 200)
        
        contacts = Contact.objects.filter(organization=self.org, email="dendroenergiapromotor@gmail.com")
        self.assertEqual(contacts.count(), 1)
        self.assertEqual(contacts.first().first_name, "Luis")

    def test_import_semicolon_csv(self):
        """Prueba importación de archivo CSV con delimitador punto y coma común en Excel en español."""
        csv_data = "email;nombre;empresa\nmaria@test.com;María;AgroChile\ncarlos@test.com;Carlos;Maderas SpA\n".encode('utf-8-sig')
        csv_file = SimpleUploadedFile("contactos_excel.csv", csv_data, content_type="text/csv")
        
        self.client.force_login(self.editor_user)
        response = self.client.post('/contacts/import/', {'file': csv_file}, follow=True)
        self.assertEqual(response.status_code, 200)
        
        self.assertEqual(Contact.objects.filter(organization=self.org).count(), 2)
        maria = Contact.objects.get(organization=self.org, email="maria@test.com")
        self.assertEqual(maria.first_name, "María")
        self.assertEqual(maria.company, "AgroChile")

    def test_import_respects_suppressions(self):
        """Verifica que contactos desuscritos previamente no sean reactivados."""
        SuppressionEntry.objects.create(
            organization=self.org,
            email="baja@test.com",
            reason="unsub"
        )
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["email", "nombre"])
        ws.append(["baja@test.com", "Persona Dada de Baja"])
        ws.append(["activa@test.com", "Persona Activa"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        excel_file = SimpleUploadedFile("supresion.xlsx", buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.client.force_login(self.admin_user)
        response = self.client.post('/contacts/import/', {'file': excel_file}, follow=True)
        self.assertEqual(response.status_code, 200)
        
        self.assertFalse(Contact.objects.filter(organization=self.org, email="baja@test.com").exists())
        self.assertTrue(Contact.objects.filter(organization=self.org, email="activa@test.com").exists())

    def test_campaign_editor_can_access_import(self):
        """Verifica que el rol campaign_editor pueda acceder e importar contactos."""
        self.client.force_login(self.editor_user)
        response = self.client.get('/contacts/import/')
        self.assertEqual(response.status_code, 200)

    def test_contact_list_pagination(self):
        """Verifica que la lista de contactos pagine adecuadamente y respete el page_size."""
        # Crear 35 contactos para la organización
        contacts = [
            Contact(organization=self.org, email=f"contacto{i}@test.com", first_name=f"User {i}")
            for i in range(35)
        ]
        Contact.objects.bulk_create(contacts)
        
        self.client.force_login(self.admin_user)
        # Default page_size es 50 -> deben mostrarse los 35 en una sola página
        resp = self.client.get('/contacts/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.context['page_obj']), 35)
        self.assertEqual(resp.context['total_count'], 35)

        # Con page_size=25 -> página 1 debe tener 25, y página 2 debe tener 10
        resp_p1 = self.client.get('/contacts/?page_size=25&page=1')
        self.assertEqual(len(resp_p1.context['page_obj']), 25)
        self.assertTrue(resp_p1.context['page_obj'].has_next())

        resp_p2 = self.client.get('/contacts/?page_size=25&page=2')
        self.assertEqual(len(resp_p2.context['page_obj']), 10)
        self.assertTrue(resp_p2.context['page_obj'].has_previous())

    def test_tag_crud_lifecycle(self):
        """Prueba ciclo completo de vida de etiquetas: crear, listar, renombrar y eliminar."""
        self.client.force_login(self.admin_user)
        
        # 1. Crear etiqueta
        resp = self.client.post('/contacts/tags/create/', {'name': 'VIP Clientes'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        tag = Tag.objects.filter(organization=self.org, name='VIP Clientes').first()
        self.assertIsNotNone(tag)
        
        # 2. Listar etiquetas
        resp_list = self.client.get('/contacts/tags/')
        self.assertEqual(resp_list.status_code, 200)
        self.assertContains(resp_list, 'VIP Clientes')
        
        # 3. Renombrar etiqueta
        resp_edit = self.client.post(f'/contacts/tags/{tag.id}/edit/', {'name': 'Super VIP'}, follow=True)
        self.assertEqual(resp_edit.status_code, 200)
        tag.refresh_from_db()
        self.assertEqual(tag.name, 'Super VIP')
        
        # 4. Eliminar etiqueta
        resp_del = self.client.post(f'/contacts/tags/{tag.id}/delete/', follow=True)
        self.assertEqual(resp_del.status_code, 200)
        self.assertFalse(Tag.objects.filter(id=tag.id).exists())

    def test_import_with_pre_created_tag(self):
        """Verifica que al importar contactos se pueda seleccionar un tag pre-creado por ID."""
        tag = Tag.objects.create(organization=self.org, name="Segmento Agro")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["email", "nombre"])
        ws.append(["agro1@test.com", "Agro Uno"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        f = SimpleUploadedFile("agro.xlsx", buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.client.force_login(self.admin_user)
        resp = self.client.post('/contacts/import/', {'file': f, 'existing_tag_id': tag.id}, follow=True)
        self.assertEqual(resp.status_code, 200)
        
        c = Contact.objects.filter(organization=self.org, email="agro1@test.com").first()
        self.assertIsNotNone(c)
        self.assertTrue(c.tags.filter(id=tag.id).exists())

    def test_contact_create_and_edit_with_tags(self):
        """Verifica selección de etiquetas existentes y adición de nuevas al crear y editar contactos."""
        tag1 = Tag.objects.create(organization=self.org, name="Tag 1")
        tag2 = Tag.objects.create(organization=self.org, name="Tag 2")
        
        self.client.force_login(self.admin_user)
        # 1. Crear contacto seleccionando tag1 y escribiendo nueva etiqueta 'Tag 3'
        resp_create = self.client.post('/contacts/create/', {
            'email': 'pepe@test.com',
            'first_name': 'Pepe',
            'tags': [str(tag1.id)],
            'new_tags': 'Tag 3',
        }, follow=True)
        self.assertEqual(resp_create.status_code, 200)
        
        c = Contact.objects.get(organization=self.org, email="pepe@test.com")
        self.assertTrue(c.tags.filter(name="Tag 1").exists())
        self.assertTrue(c.tags.filter(name="Tag 3").exists())
        self.assertFalse(c.tags.filter(name="Tag 2").exists())
        
        # 2. Editar contacto cambiando a tag2 y agregando 'Tag 4'
        tag3 = Tag.objects.get(organization=self.org, name="Tag 3")
        resp_edit = self.client.post(f'/contacts/{c.id}/edit/', {
            'email': 'pepe@test.com',
            'first_name': 'Pepe Editado',
            'tags': [str(tag2.id)],
            'new_tags': 'Tag 4',
        }, follow=True)
        self.assertEqual(resp_edit.status_code, 200)
        
        c.refresh_from_db()
        self.assertEqual(c.first_name, 'Pepe Editado')
        self.assertFalse(c.tags.filter(name="Tag 1").exists())
        self.assertTrue(c.tags.filter(name="Tag 2").exists())
        self.assertTrue(c.tags.filter(name="Tag 4").exists())


